"""
excel.py  --  Turn grades into an .xlsx file, and an .xlsx file back into grades.
==================================================================================

We use the `openpyxl` library. The exported workbook has 3 sheets:

  "Summary"  -> for READING: each course's current %, grade and the GWA.
  "Data"     -> for RESTORING: one row per score. This is the sheet that is
                read back when you import. You may edit it in Excel too!
  "Settings" -> passing percent and the grade scale.

Because "Data" is a plain table, the exported file works as a "save file":
if your browser's local storage ever gets wiped, import the .xlsx to get
everything back.
"""

from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .grading import component_average, course_standing, gwa
from .models import (
    MAX_COURSES, AppData, Component, Course, GradeRow, Item, Settings,
    clean_number, clean_text,
)

# Forest green header, to match the app's dark theme.
HEADER_FILL = PatternFill("solid", fgColor="1F6B35")
HEADER_FONT = Font(bold=True, color="FFFFFF")
MAX_IMPORT_ROWS = 5000

DATA_HEADERS = ["Course", "Code", "Units", "Component", "Weight %", "Item", "Score", "Max", "Percent"]


# ---------------------------------------------------------------------------
# Small helpers
# ---------------------------------------------------------------------------
def _put(ws, row: int, col: int, value):
    """
    Write a value into a cell. Text is forced to be plain text so that a name
    like '=1+1' can never run as an Excel formula (a known spreadsheet trick).
    """
    cell = ws.cell(row=row, column=col)
    cell.value = value
    if isinstance(value, str):
        cell.data_type = "s"
    return cell


def _style_header(ws, headers: list[str]):
    for col, title in enumerate(headers, start=1):
        cell = _put(ws, 1, col, title)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "A2"


def _widths(ws, widths: list[int]):
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width


# ---------------------------------------------------------------------------
# EXPORT: AppData -> .xlsx bytes
# ---------------------------------------------------------------------------
def build_workbook(data: AppData) -> BytesIO:
    wb = Workbook()

    # ---- Sheet 1: Summary ----
    ws = wb.active
    ws.title = "Summary"
    _style_header(ws, ["Course", "Code", "Units", "Current %", "Grade", "Status", "Graded weight %"])
    row = 2
    for course in data.courses:
        s = course_standing(course, data.settings)
        _put(ws, row, 1, course.name)
        _put(ws, row, 2, course.code)
        _put(ws, row, 3, course.units)
        _put(ws, row, 4, round(s["percent"], 2) if s["percent"] is not None else "")
        _put(ws, row, 5, s["grade"] if s["grade"] is not None else "")
        status = "No scores yet" if s["passed"] is None else ("Passing" if s["passed"] else "Below passing")
        _put(ws, row, 6, status)
        _put(ws, row, 7, round(s["graded_weight"], 2))
        row += 1
    overall = gwa(data.courses, data.settings)
    _put(ws, row + 1, 1, "GWA").font = Font(bold=True)
    _put(ws, row + 1, 5, round(overall, 4) if overall is not None else "").font = Font(bold=True)
    _widths(ws, [34, 12, 8, 12, 8, 16, 16])

    # ---- Sheet 2: Data (this is what Import reads) ----
    wd = wb.create_sheet("Data")
    _style_header(wd, DATA_HEADERS)
    row = 2
    for course in data.courses:
        comps = course.components or [Component(name="", weight=0)]
        for comp in comps:
            items = comp.items or [None]  # keep empty components in the file too
            for item in items:
                _put(wd, row, 1, course.name)
                _put(wd, row, 2, course.code)
                _put(wd, row, 3, course.units)
                _put(wd, row, 4, comp.name)
                _put(wd, row, 5, comp.weight)
                if item is not None:
                    _put(wd, row, 6, item.name)
                    _put(wd, row, 7, item.score if item.score is not None else "")
                    _put(wd, row, 8, item.max_score)
                    # A live Excel formula so the % updates if you edit scores in Excel.
                    pct = wd.cell(row=row, column=9)
                    pct.value = f'=IF(AND(ISNUMBER(G{row}),ISNUMBER(H{row}),H{row}>0),G{row}/H{row},"")'
                    pct.number_format = "0.0%"
                row += 1
    _widths(wd, [34, 12, 8, 22, 10, 24, 10, 10, 10])

    # ---- Sheet 3: Settings ----
    wt = wb.create_sheet("Settings")
    _style_header(wt, ["Setting", "Value"])
    s = data.settings
    for i, (label, value) in enumerate([
        ("Passing percent", s.passing_percent),
        ("Failing grade", s.fail_grade),
        ("Round percent", "yes" if s.round_percent else "no"),
    ], start=2):
        _put(wt, i, 1, label)
        _put(wt, i, 2, value)
    _put(wt, 6, 1, "Grade scale: from percent").font = Font(bold=True)
    _put(wt, 6, 2, "Grade").font = Font(bold=True)
    for i, r in enumerate(s.scale, start=7):
        _put(wt, i, 1, r.min_percent)
        _put(wt, i, 2, r.grade)
    _widths(wt, [28, 14])

    out = BytesIO()
    wb.save(out)
    out.seek(0)
    return out


# ---------------------------------------------------------------------------
# IMPORT: .xlsx file -> AppData
# ---------------------------------------------------------------------------
def _find_sheet(wb, name: str):
    for title in wb.sheetnames:
        if title.strip().lower() == name.lower():
            return wb[title]
    return None


def parse_workbook(file_like) -> AppData:
    """Read an .xlsx made by build_workbook() (or edited by hand) back into AppData."""
    try:
        wb = load_workbook(file_like, data_only=True, read_only=True)
    except Exception as exc:  # openpyxl raises many different errors for bad files
        raise ValueError("That file could not be read as an Excel (.xlsx) workbook.") from exc

    sheet = _find_sheet(wb, "Data")
    if sheet is None:
        raise ValueError('This workbook has no "Data" sheet. Use a file exported from this app.')

    courses: dict[tuple, Course] = {}
    comps: dict[tuple, Component] = {}
    for n, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        if n > MAX_IMPORT_ROWS:
            raise ValueError(f"Too many rows (limit {MAX_IMPORT_ROWS}).")
        row = list(row) + [None] * (9 - len(row))
        course_name, code, units, comp_name, weight, item_name, score, max_score = row[:8]
        if not clean_text(course_name):
            continue  # blank line

        ckey = (clean_text(course_name).lower(), clean_text(code, "", 30).lower())
        if ckey not in courses:
            if len(courses) >= MAX_COURSES:
                raise ValueError(f"Too many courses (limit {MAX_COURSES}).")
            courses[ckey] = Course(
                name=clean_text(course_name), code=clean_text(code, "", 30),
                units=clean_number(units, 3.0, 0, 30),
            )
        course = courses[ckey]

        if clean_text(comp_name):
            comp_key = (ckey, clean_text(comp_name).lower())
            if comp_key not in comps:
                comps[comp_key] = Component(name=clean_text(comp_name),
                                            weight=clean_number(weight, 0.0, 0, 100))
                course.components.append(comps[comp_key])
            comp = comps[comp_key]
            if clean_text(item_name) or score not in (None, ""):
                comp.items.append(Item(
                    name=clean_text(item_name, f"Item {len(comp.items) + 1}"),
                    score=clean_number(score, None, 0),
                    max_score=clean_number(max_score, 100.0, 0.0001) or 100.0,
                ))

    settings = Settings()
    sheet = _find_sheet(wb, "Settings")
    if sheet is not None:
        rows = [list(r) + [None] * 2 for r in sheet.iter_rows(min_row=2, values_only=True)]
        scale: list[GradeRow] = []
        in_scale = False
        for a, b, *_ in rows:
            label = clean_text(a).lower()
            if label.startswith("grade scale"):
                in_scale = True
            elif in_scale:
                pct, grade = clean_number(a, None, 0, 100), clean_number(b, None, 0, 100)
                if pct is not None and grade is not None:
                    scale.append(GradeRow(pct, grade))
            elif label == "passing percent":
                settings.passing_percent = clean_number(b, 75.0, 0, 100)
            elif label == "failing grade":
                settings.fail_grade = clean_number(b, 5.0, 0, 100)
            elif label == "round percent":
                settings.round_percent = clean_text(b).lower() not in ("no", "false", "0")
        if scale:
            settings.scale = sorted(scale, key=lambda r: r.min_percent, reverse=True)[:40]

    wb.close()
    if not courses:
        raise ValueError("No courses were found in the Data sheet.")
    return AppData(settings=settings, courses=list(courses.values()))
